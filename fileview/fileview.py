#!/usr/bin/env python3
"""Tailnet file viewer: one link per file, resolved by who clicks it.

URL path == filesystem path, e.g.
    https://my-machine.tailnet.ts.net/home/me/project/README.md:12

Every client gets a browser IDE (file tree, tabs, file search, outline). Clients
listed under "editors" in the config also get "Open in <editor>" buttons. Local
clients jump straight into VS Code (text files, folders) or the desktop app
(office files, PDFs).

Machine-specific settings (public URL, machine name, local clients, editor links)
come from ~/.config/fileview/config.json, or $FILEVIEW_CONFIG; see
config.example.json. Without one, fileview serves localhost only.

Query flags: ?view (stay in the browser), ?src (source of md/html), ?raw (bytes),
?embed (html page itself), ?ls (directory entries as JSON), ?find=q (file search),
?pages (page images of a .docx), ?page=N (one page or slide as PNG),
?media=ppt/media/x.mp4 (a video embedded in a .pptx, overlaid on its slide).
Slides and pages come from LibreOffice -> PDF -> pdftoppm, cached in ~/.cache/fileview.
The X-FV-Partial header returns a page as JSON for in-app navigation.
UI assets live in static/ next to this script and are served at /_fv/static/.
"""
import csv
import datetime
import gzip
import hashlib
import heapq
import html
import io
import json
import mimetypes
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import urllib.parse
import xml.etree.ElementTree as ET
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import mammoth
import markdown
import openpyxl
from pygments import highlight
from pygments.formatters import HtmlFormatter
from pygments.lexers import TextLexer, get_lexer_by_name, get_lexer_for_filename
from pygments.util import ClassNotFound

HOST = "127.0.0.1"
PORT = int(os.environ.get("FILEVIEW_PORT", "8765"))
HOME = os.path.expanduser("~")
ROOTS = [os.path.realpath(p) for p in (HOME, "/tmp", "/opt")]
STATIC = os.path.join(os.path.dirname(os.path.realpath(__file__)), "static")
CACHE = os.path.expanduser("~/.cache/fileview")

CONFIG_PATH = os.environ.get("FILEVIEW_CONFIG", os.path.expanduser("~/.config/fileview/config.json"))


def load_config():
    try:
        with open(CONFIG_PATH) as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


CONFIG = load_config()
PUBLIC = CONFIG.get("public_url", f"http://{HOST}:{PORT}").rstrip("/")
MACHINE = CONFIG.get("machine", socket.gethostname())
# Clients on this machine: links jump straight into local VS Code or the desktop app.
LOCAL_CLIENTS = {"127.0.0.1", *CONFIG.get("local_clients", [])}
LOCAL_VSCODE = [("VS Code", "vscode://file{path}:{line}")]
# client IP -> [(label, URL template with {path}, {line}, {end})]
EDITORS = {ip: LOCAL_VSCODE for ip in LOCAL_CLIENTS}
EDITORS.update({ip: [tuple(e) for e in editors] for ip, editors in CONFIG.get("editors", {}).items()})

DENY_DIRS = {".ssh", ".gnupg", ".aws", ".azure", ".kube", ".docker", ".password-store", "keyrings"}
DENY_FILE = re.compile(
    r"^(\.env(\..+)?|.*\.pem|.*\.key|id_(rsa|ed25519|ecdsa)(\.pub)?|\.netrc|\.git-credentials"
    r"|\.?credentials(\.json)?|auth\.json|hosts\.yml)$"
)
ALLOW_FILE = re.compile(r".*\.(example|sample|template)$")
TREE_HIDE = {".git", "__pycache__", ".DS_Store"}
TREE_MAX = 2000
INDEX_PRUNE = TREE_HIDE | {"node_modules", "build", "devel", "logs", "install", ".cache", ".venv", "venv",
                           "dist", "target", ".npm", ".cargo", ".rustup", ".bun", ".nvm", ".mozilla",
                           ".vscode-server", ".cursor-server", ".t3", ".local", ".catkin_tools"}
INDEX_MAX = 300_000
INDEX_TTL = 120
FIND_LIMIT = 60

MAX_HIGHLIGHT = 1_500_000
MAX_TEXT = 5_000_000
LOG_TAIL_LINES = 3000
CSV_ROWS = 1000
SHEET_ROWS = 2000
SHEET_COLS = 200
PAGE_PX = 1600  # long side of rendered slides and pages
CONVERT_TIMEOUT = 180
LINE_SUFFIX = re.compile(r"^(.*?):(\d+)(?:-(\d+))?(?::\d+)?$")
LEXER_ALIASES = {".launch": "xml", ".test": "xml", ".xacro": "xml", ".urdf": "xml", ".world": "xml",
                 ".beh": "json", ".pml": "json", ".repos": "yaml", ".rosinstall": "yaml",
                 ".msg": "text", ".srv": "text", ".action": "text", "Tiltfile": "python"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".bmp"}
VIDEO_EXTS = {".mp4", ".m4v", ".webm", ".mov", ".ogv"}
mimetypes.add_type("video/mp4", ".m4v")
DESKTOP_EXTS = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".odt", ".ods", ".odp"}
DOWNLOAD_EXTS = DESKTOP_EXTS | {".zip", ".bag", ".db3", ".mcap"}
BROWSER_EXTS = {".html", ".htm", ".csv", ".tsv"} | DOWNLOAD_EXTS | VIDEO_EXTS
SLIDE_EXTS = {".pptx", ".odp"}
SHEET_EXTS = {".xlsx", ".xlsm"}

FILE_TYPES = {
    "t-md": ".md .markdown .rst .txt",
    "t-cpp": ".cpp .cc .c .h .hpp .js .mjs .ts .tsx .go .rs .java",
    "t-py": ".py .sh .bash .zsh Tiltfile Makefile CMakeLists.txt Dockerfile",
    "t-data": ".json .yaml .yml .xml .toml .launch .test .xacro .urdf .beh .pml .csv .tsv .repos",
    "t-web": ".html .htm .css",
    "t-img": " ".join(sorted(IMAGE_EXTS | VIDEO_EXTS)),
    "t-office": " ".join(sorted(DOWNLOAD_EXTS)),
    "t-log": ".log .out .gz",
}
TYPE_OF = {key: cls for cls, keys in FILE_TYPES.items() for key in keys.split()}


def file_type(name):
    ext = os.path.splitext(name)[1].lower()
    return TYPE_OF.get(name) or TYPE_OF.get(ext, "")


def is_denied(real):
    parts = real.split(os.sep)
    if any(p in DENY_DIRS for p in parts):
        return True
    if "/.config/gh" in real or real.endswith("/.claude/.credentials.json"):
        return True
    name = parts[-1]
    return bool(DENY_FILE.match(name)) and not ALLOW_FILE.match(name)


def inside_roots(real):
    return any(real == r or real.startswith(r + os.sep) for r in ROOTS)


def tree_root(path):
    """Topmost git checkout above path, else the top-level folder under its root."""
    d = (path if path.endswith("/") or os.path.isdir(path) else os.path.dirname(path)).rstrip("/") or "/"
    base = next((r for r in ROOTS if d == r or d.startswith(r + "/")), None)
    if base is None or d == base:
        return d.rstrip("/") + "/"
    top, cur = None, d
    while cur != base:
        if os.path.exists(os.path.join(cur, ".git")):
            top = cur
        cur = os.path.dirname(cur)
    return (top or base + "/" + d[len(base) + 1:].split("/")[0]) + "/"


def natural(name):
    return [(0, int(t), "") if t.isdigit() else (1, 0, t) for t in re.split(r"(\d+)", name.lower()) if t]


def git_info(path):
    """Read the containing checkout's HEAD, including nested repos and worktrees."""
    real = os.path.realpath(path)
    if not inside_roots(real) or is_denied(real) or not os.path.exists(real):
        return None
    directory = real if os.path.isdir(real) else os.path.dirname(real)
    try:
        def git(*args):
            result = subprocess.run(["git", "-C", directory] + list(args),
                                    stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                    encoding="utf-8", errors="replace", timeout=2)
            return result.stdout.strip() if result.returncode == 0 else ""

        branch = git("symbolic-ref", "--quiet", "--short", "HEAD")
        if branch:
            return {"branch": branch, "detached": False}
        commit = git("rev-parse", "--short", "HEAD")
        if commit:
            return {"branch": commit, "detached": True}
    except (OSError, subprocess.TimeoutExpired):
        pass
    return None


def human(n):
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


def opens_in_editor(real):
    ext = os.path.splitext(real)[1].lower()
    return ext not in IMAGE_EXTS | BROWSER_EXTS and not real.endswith(".gz")


def lexer_for(path, text):
    name = os.path.basename(path)
    if name.endswith(".gz"):
        name = name[:-3]
    ext = os.path.splitext(name)[1].lower()
    alias = LEXER_ALIASES.get(name) or LEXER_ALIASES.get(ext)
    try:
        if alias:
            return get_lexer_by_name(alias)
        return get_lexer_for_filename(name, text[:4096])
    except ClassNotFound:
        return TextLexer()


def read_text(path):
    """Returns (text, note). Handles .gz, big logs and binary detection."""
    size = os.path.getsize(path)
    opener = gzip.open if path.endswith(".gz") else open
    ext = os.path.splitext(path[:-3] if path.endswith(".gz") else path)[1].lower()
    with opener(path, "rb") as f:
        if ext in (".log", ".out") and size > MAX_TEXT and not path.endswith(".gz"):
            f.seek(max(0, size - MAX_TEXT))
            data = f.read()
            lines = data.decode("utf-8", "replace").splitlines()[-LOG_TAIL_LINES:]
            return "\n".join(lines), f"Large log ({human(size)}): showing the last {len(lines)} lines."
        data = f.read(MAX_TEXT + 1)
    note = ""
    if len(data) > MAX_TEXT:
        data = data[:MAX_TEXT]
        note = f"Truncated to the first {human(MAX_TEXT)}."
    if b"\x00" in data[:8192]:
        return None, "Binary file."
    text = data.decode("utf-8", "replace")
    if ext == ".json" and len(text) < MAX_HIGHLIGHT:
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            pass
    return text, note


def render_markdown(text):
    return markdown.markdown(text, extensions=["extra", "sane_lists", "toc", "codehilite"],
                             extension_configs={"codehilite": {"css_class": "hl", "guess_lang": False}})


_convert_lock = threading.Lock()


def page_images(real):
    """PNG paths, one per slide or page; converted once per file version and cached."""
    st = os.stat(real)
    out = os.path.join(CACHE, hashlib.sha1(real.encode()).hexdigest()[:16])
    stamp = f"{st.st_mtime_ns} {st.st_size}"
    done = os.path.join(out, "stamp")

    def cached():
        try:
            with open(done) as f:
                if f.read() == stamp:
                    return sorted((os.path.join(out, n) for n in os.listdir(out) if n.endswith(".png")),
                                  key=natural)
        except OSError:
            pass
        return None

    with _convert_lock:  # one LibreOffice at a time; it shares the profile below
        hit = cached()
        if hit is not None:
            return hit
        shutil.rmtree(out, ignore_errors=True)
        os.makedirs(out)
        with tempfile.TemporaryDirectory() as tmp:
            pdf = real
            if not real.lower().endswith(".pdf"):
                subprocess.run(["soffice", f"-env:UserInstallation=file://{CACHE}/lo-profile", "--headless",
                                "--convert-to", "pdf", "--outdir", tmp, real],
                               stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                               timeout=CONVERT_TIMEOUT, check=True)
                pdf = os.path.join(tmp, os.path.splitext(os.path.basename(real))[0] + ".pdf")
            subprocess.run(["pdftoppm", "-png", "-scale-to", str(PAGE_PX), pdf, os.path.join(out, "p")],
                           stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE,
                           timeout=CONVERT_TIMEOUT, check=True)
        with open(done, "w") as f:
            f.write(stamp)
        return cached()


PPTX_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
           "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
           "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
           "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
           "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main"}


def slide_videos(real):
    """Embedded videos per shown slide: [[{part, x, y, w, h}], ...] with the box as % of the slide."""
    def rels(z, part):
        folder, name = os.path.split(part)
        try:
            root = ET.fromstring(z.read(f"{folder}/_rels/{name}.rels"))
        except KeyError:
            return {}
        return {r.get("Id"): os.path.normpath(os.path.join(folder, r.get("Target")))
                for r in root.iterfind("pr:Relationship", PPTX_NS) if r.get("TargetMode") != "External"}

    with zipfile.ZipFile(real) as z:
        pres = ET.fromstring(z.read("ppt/presentation.xml"))
        size = pres.find("p:sldSz", PPTX_NS)
        sw, sh = int(size.get("cx")), int(size.get("cy"))
        pres_rels = rels(z, "ppt/presentation.xml")
        out = []
        for sid in pres.iterfind("p:sldIdLst/p:sldId", PPTX_NS):
            part = pres_rels.get(sid.get(f"{{{PPTX_NS['r']}}}id"))
            if not part:
                continue
            slide = ET.fromstring(z.read(part))
            if slide.get("show") == "0":  # hidden slides are left out of the export
                continue
            slide_rels, videos = rels(z, part), []
            for pic in slide.iterfind("p:cSld/p:spTree/p:pic", PPTX_NS):
                ref = pic.find(".//p14:media", PPTX_NS)
                if ref is None:
                    ref = pic.find(".//a:videoFile", PPTX_NS)
                media = None if ref is None else slide_rels.get(
                    ref.get(f"{{{PPTX_NS['r']}}}embed") or ref.get(f"{{{PPTX_NS['r']}}}link"))
                off, ext = pic.find("p:spPr/a:xfrm/a:off", PPTX_NS), pic.find("p:spPr/a:xfrm/a:ext", PPTX_NS)
                if media in z.namelist() and off is not None and ext is not None:
                    videos.append({"part": media, "x": int(off.get("x")) / sw * 100, "y": int(off.get("y")) / sh * 100,
                                   "w": int(ext.get("cx")) / sw * 100, "h": int(ext.get("cy")) / sh * 100})
            out.append(videos)
        return out


def media_part(real, part):
    """Extract one embedded media file to the cache (per file version) and return its path."""
    st = os.stat(real)
    key = hashlib.sha1(f"{real} {st.st_mtime_ns} {st.st_size} {part}".encode()).hexdigest()[:16]
    dest = os.path.join(CACHE, "media", key + os.path.splitext(part)[1].lower())
    if not os.path.exists(dest):
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        with zipfile.ZipFile(real) as z, z.open(part) as src, tempfile.NamedTemporaryFile(
                dir=os.path.dirname(dest), delete=False) as tmp:
            shutil.copyfileobj(src, tmp)
        os.replace(tmp.name, dest)
    return dest


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def cell_text(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time() else v.isoformat(" ", "seconds")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, (int, float)):
        fmt = cell.number_format or "General"
        m = re.search(r"0\.(0+)", fmt)
        places = len(m.group(1)) if m else None
        if "%" in fmt:
            return f"{v * 100:.{places or 0}f}%"
        sep = "," if "#,##" in fmt else ""
        if places is not None:
            return f"{v:{sep}.{places}f}"
        if float(v).is_integer():
            return f"{int(v):{sep}d}"
        return f"{v:.10g}"
    return str(v)


def argb(color):
    rgb = getattr(color, "rgb", None) if color is not None and getattr(color, "type", None) == "rgb" else None
    return f"#{rgb[-6:]}" if isinstance(rgb, str) and len(rgb) >= 6 and rgb[-6:] != "000000" else None


def cell_style(cell):
    css = []
    fill = cell.fill
    bg = argb(fill.fgColor) if fill is not None and fill.fill_type == "solid" else None
    fg = argb(cell.font.color) if cell.font is not None else None
    if bg and bg.upper() != "#FFFFFF":
        css += [f"background:{bg}", f"color:{fg or '#000'}"]
    elif fg:
        css.append(f"color:{fg}")
    if cell.font is not None:
        if cell.font.b:
            css.append("font-weight:600")
        if cell.font.i:
            css.append("font-style:italic")
    align = cell.alignment.horizontal if cell.alignment is not None else None
    if align in ("left", "center", "right"):
        css.append(f"text-align:{align}")
    return ";".join(css)


def sheet_table(ws):
    """One worksheet as an HTML table with spreadsheet-style row and column headers."""
    rows = [r for r in ws.iter_rows(max_row=min(ws.max_row, SHEET_ROWS), max_col=min(ws.max_column, SHEET_COLS))]
    last_row = max((i for i, r in enumerate(rows, 1) if any(c.value is not None for c in r)), default=0)
    last_col = max((c.column for r in rows[:last_row] for c in r if c.value is not None), default=0)
    if not last_row:
        return '<div class="fnone">This sheet is empty.</div>', ""
    spans, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        spans[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        covered.update((r, c) for r in range(rng.min_row, rng.max_row + 1)
                       for c in range(rng.min_col, rng.max_col + 1) if (r, c) != (rng.min_row, rng.min_col))
    hidden_cols = {c for c in range(1, last_col + 1)
                   if ws.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden}
    cols = [c for c in range(1, last_col + 1) if c not in hidden_cols]
    out = ['<thead><tr><th class="corner"></th>' +
           "".join(f"<th>{openpyxl.utils.get_column_letter(c)}</th>" for c in cols) + "</tr></thead><tbody>"]
    for r, row in enumerate(rows[:last_row], 1):
        if ws.row_dimensions[r].hidden:
            continue
        tds = []
        for cell in row[:last_col]:
            c = cell.column
            if c in hidden_cols or (r, c) in covered:
                continue
            rs, cs = spans.get((r, c), (1, 1))
            attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                attrs += ' class="n"'
            style = cell_style(cell)
            if style:
                attrs += f' style="{style}"'
            tds.append(f"<td{attrs}>{html.escape(cell_text(cell))}</td>")
        out.append(f'<tr><th class="rn">{r}</th>{"".join(tds)}</tr>')
    out.append("</tbody>")
    clipped = ws.max_row > SHEET_ROWS or ws.max_column > SHEET_COLS
    note = f"Showing the first {SHEET_ROWS} rows and {SHEET_COLS} columns." if clipped else ""
    return f'<div class="grid sheet-grid"><table>{"".join(out)}</table></div>', note


_index, _index_lock = {}, threading.Lock()


def file_index(root):
    with _index_lock:
        hit = _index.get(root)
        if hit and time.time() - hit[0] < INDEX_TTL:
            return hit[1]
        files = []
        for d, dirs, names in os.walk(root):
            dirs[:] = [x for x in dirs if x not in INDEX_PRUNE and not is_denied(os.path.join(d, x))]
            for n in names:
                p = os.path.join(d, n)
                if not is_denied(p):
                    rel = p[len(root):]
                    files.append((rel, rel.lower(), n.lower()))
            if len(files) >= INDEX_MAX:
                break
        _index[root] = (time.time(), files)
        return files


def fuzzy_score(text, q):
    """Tightest subsequence window of q in text; matches at word starts count as closer."""
    best = None
    for i, c in enumerate(text):
        if c != q[0]:
            continue
        j, k, starts = i, 0, 0
        while j < len(text) and k < len(q):
            if text[j] == q[k]:
                starts += j == 0 or not text[j - 1].isalnum()
                k += 1
            j += 1
        if k < len(q):
            break
        if j - i <= 3 * len(q):
            score = (j - i) / len(q) - starts / 4
            best = score if best is None else min(best, score)
    return best


def search(files, query):
    q = re.sub(r"\s+", "", query.lower())
    if not q:
        return []
    # [^a]*a[^b]*b... is a linear prefilter for the subsequence check
    rx = re.compile("".join(f"[^{c}]*{c}" for c in map(re.escape, q)))
    exact, fuzzy = [], []
    for rel, low, base in files:
        if q in base:
            exact.append(((0 if base.startswith(q) else 1) + len(low) / 1000, rel))
        elif q in low:
            exact.append((3 + len(low) / 1000, rel))
        elif not exact:
            text = low if "/" in q else base
            score = rx.match(text) and fuzzy_score(text, q)
            if score is not None and score is not False:
                fuzzy.append((score + len(low) / 1000, rel))
    return [rel for _, rel in heapq.nsmallest(FIND_LIMIT, exact or fuzzy)]


def asset(name):
    return f"/_fv/static/{name}?v={int(os.path.getmtime(os.path.join(STATIC, name)))}"


def empty_state(icon, title, text, buttons=()):
    links = "".join(f'<a class="btn{" primary" if primary else ""}" href="{html.escape(href)}"{attrs}>{html.escape(label)}</a>'
                    for label, href, primary, attrs in buttons)
    return (f'<div class="empty"><div class="empty-icon" data-icon="{icon}"></div><h2>{html.escape(title)}</h2>'
            f'<p>{text}</p><div class="empty-actions">{links}</div></div>')


HEAD_JS = """<script>
try { var d = document.documentElement, t = localStorage.getItem('fv.theme'), w = localStorage.getItem('fv.w');
  if (t === '"light"' || t === '"dark"') d.dataset.theme = JSON.parse(t);
  if (w) d.style.setProperty('--side-w', w + 'px');
  if (localStorage.getItem('fv.hidden') === 'true') d.classList.add('nav-hidden');
  if (localStorage.getItem('fv.wrap') === 'true') d.classList.add('wrap'); } catch (e) {}
</script>"""


def shell(parts, body):
    state = json.dumps(parts).replace("</", "<\\/")
    title = html.escape(parts["title"])
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>{title}</title>
<link rel="preload" href="/_fv/static/Geist.woff2" as="font" type="font/woff2" crossorigin>
<link rel="preload" href="/_fv/static/GeistMono.woff2" as="font" type="font/woff2" crossorigin>
<link rel="stylesheet" href="{asset('app.css')}">{HEAD_JS}</head><body>
<div class="app" id="app">
  <header class="top" id="top"></header>
  <div class="frame">
    <aside class="side" id="side"><div class="side-head" id="sidehead"></div><nav class="tree" id="tree"></nav><div class="grip" id="grip"></div></aside>
    <div class="scrim" id="scrim"></div>
    <section class="editor">
      <div class="tabs" id="tabs"></div>
      <div class="crumbbar"><nav class="crumbs" id="crumbs"></nav><div class="actions" id="actions"></div></div>
      <div class="progress" id="progress"></div>
      <div class="content" id="content"><main id="main" class="{parts["kind"]}">{body}</main><aside class="toc" id="toc"></aside></div>
    </section>
  </div>
  <footer class="status" id="status"></footer>
</div>
<script type="application/json" id="fv">{state}</script>
<script src="{asset('app.js')}"></script>
</body></html>"""


class Handler(BaseHTTPRequestHandler):
    server_version = "fileview/2"
    partial = False

    def log_message(self, fmt, *args):
        sys.stderr.write("%s %s\n" % (self.client_ip(), fmt % args))

    def client_ip(self):
        fwd = self.headers.get("X-Forwarded-For")
        return fwd.split(",")[0].strip() if fwd else self.client_address[0]

    def send(self, code, body, ctype="text/html; charset=utf-8", headers=None):
        data = body.encode() if isinstance(body, str) else body
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()
        if self.command != "HEAD":
            self.wfile.write(data)

    def send_json(self, code, obj):
        self.send(code, json.dumps(obj), "application/json")

    def respond(self, code, path, body, kind="doc", info=None, actions=(), line=None, end=None, transient=False):
        is_dir = path.endswith("/")
        parts = {
            "path": path, "title": os.path.basename(path.rstrip("/")) or "/", "kind": kind, "isDir": is_dir,
            "root": tree_root(path), "info": info or {}, "actions": list(actions), "line": line, "end": end,
            "editors": [] if is_dir or code != 200 else [
                {"label": label, "tpl": tpl.replace("{path}", urllib.parse.quote(path))}
                for label, tpl in EDITORS.get(self.client_ip(), [])],
            "transient": transient, "public": PUBLIC, "machine": MACHINE, "roots": ROOTS, "types": TYPE_OF,
            "git": git_info(path) if code == 200 else None,
        }
        hit = _index.get(parts["root"])
        if not hit or time.time() - hit[0] > INDEX_TTL:
            threading.Thread(target=file_index, args=(parts["root"],), daemon=True).start()
        if self.partial:
            parts["body"] = body
            return self.send_json(code, parts)
        self.send(code, shell(parts, body))

    def do_HEAD(self):
        self.do_GET()

    def do_GET(self):
        url = urllib.parse.urlsplit(self.path)
        if url.path.startswith("/_fv/static/"):
            return self.static(url.path[len("/_fv/static/"):])
        query = urllib.parse.parse_qs(url.query, keep_blank_values=True)
        path = urllib.parse.unquote(url.path)
        self.partial = self.headers.get("X-FV-Partial") == "1"
        if path in ("", "/"):
            return self.redirect(urllib.parse.quote(HOME) + "/")

        line = end = None
        if not os.path.lexists(path):
            m = LINE_SUFFIX.match(path)
            if m and os.path.lexists(m.group(1)):
                path, line = m.group(1), int(m.group(2))
                end = int(m.group(3)) if m.group(3) else None

        real = os.path.realpath(path)
        if not inside_roots(real) or is_denied(real):
            return self.respond(403, path, empty_state("lock", "Not allowed", "This path is outside the shared folders."),
                                kind="empty")
        if not os.path.exists(real):
            return self.respond(404, path, empty_state("search", "Not found", "No file or folder at this path."),
                                kind="empty")

        if "git" in query:
            return self.send_json(200, {"git": git_info(real)})

        local = self.client_ip() in LOCAL_CLIENTS and "view" not in query and not self.partial
        if os.path.isdir(real):
            if not url.path.endswith("/"):
                return self.redirect(urllib.parse.quote(path) + "/" + (f"?{url.query}" if url.query else ""))
            path = path.rstrip("/") + "/"
            if "ls" in query:
                return self.ls_json(real)
            if "find" in query:
                return self.send_json(200, {"results": search(file_index(path), query["find"][0])})
            if local:
                return self.jump_to_editor(path, None)
            return self.listing(path, real)

        if "raw" in query:
            return self.raw(real)
        if "media" in query and real.lower().endswith(".pptx"):
            return self.slide_media(real, query["media"][0])
        if "page" in query:
            return self.page_png(real, query["page"][0])
        if "embed" in query and real.lower().endswith((".html", ".htm")):
            return self.raw(real, "text/html; charset=utf-8")

        if local and opens_in_editor(real):
            return self.jump_to_editor(path, line)
        if local and os.path.splitext(real)[1].lower() in DESKTOP_EXTS:
            return self.open_on_desktop(path, real)

        try:
            return self.render(path, real, line, end, "src" in query, "pages" in query)
        except PermissionError:
            return self.respond(403, path, empty_state("lock", "Permission denied", "The server cannot read this file."),
                                kind="empty")

    def static(self, name):
        full = os.path.join(STATIC, os.path.basename(name))
        if not os.path.isfile(full):
            return self.send(404, "not found", "text/plain")
        ctype = mimetypes.guess_type(full)[0] or "application/octet-stream"
        with open(full, "rb") as f:
            self.send(200, f.read(), ctype, {"Cache-Control": "public, max-age=31536000, immutable"})

    def redirect(self, location):
        self.send_response(302)
        self.send_header("Location", location)
        self.send_header("Content-Length", "0")
        self.end_headers()

    def jump_to_editor(self, path, line):
        is_dir = path.endswith("/")
        target = f"vscode://file{urllib.parse.quote(path.rstrip('/'))}" + ("" if is_dir else f":{line or 1}")
        here = f"{urllib.parse.quote(path)}{f':{line}' if line else ''}?view"
        body = empty_state("code", "Opening in VS Code", "The file opens in your local VS Code window.",
                           [("Retry", target, False, ""), ("View here", here, True, "")])
        body += f"<script>location.replace({json.dumps(target)})</script>"
        self.respond(200, path, body, kind="empty", transient=True)

    def open_on_desktop(self, path, real):
        if self.command == "GET":
            subprocess.Popen(["xdg-open", real], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                             stderr=subprocess.DEVNULL, start_new_session=True)
        here = f"{urllib.parse.quote(path)}?view"
        body = empty_state("external", "Opened in the desktop app", "The file opened with its default application.",
                           [("View here", here, True, "")])
        self.respond(200, path, body, kind="empty", transient=True)

    def raw(self, real, ctype=None):
        if ctype is None:
            ctype = mimetypes.guess_type(real)[0] or "application/octet-stream"
            if ctype.startswith("text/") or ctype in ("application/json", "application/x-sh"):
                ctype = "text/plain; charset=utf-8"
        size = os.path.getsize(real)
        start, stop = 0, size - 1
        m = re.fullmatch(r"bytes=(\d*)-(\d*)", self.headers.get("Range", "").strip())
        if m and (m.group(1) or m.group(2)):
            if m.group(1):
                start, stop = int(m.group(1)), min(int(m.group(2) or stop), stop)
            else:  # suffix range: the last N bytes
                start = max(size - int(m.group(2)), 0)
            if start > stop:
                self.send_response(416)
                self.send_header("Content-Range", f"bytes */{size}")
                self.send_header("Content-Length", "0")
                return self.end_headers()
        partial = (start, stop) != (0, size - 1)
        self.send_response(206 if partial else 200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(stop - start + 1))
        self.send_header("Accept-Ranges", "bytes")
        if partial:
            self.send_header("Content-Range", f"bytes {start}-{stop}/{size}")
        self.send_header("Content-Security-Policy", "sandbox")
        self.end_headers()
        if self.command != "HEAD":
            with open(real, "rb") as f:
                f.seek(start)
                left = stop - start + 1
                try:
                    while left > 0:
                        chunk = f.read(min(left, 1 << 20))
                        if not chunk:
                            break
                        self.wfile.write(chunk)
                        left -= len(chunk)
                except (BrokenPipeError, ConnectionResetError):
                    pass  # players drop connections while seeking

    def slide_media(self, real, part):
        try:
            allowed = {v["part"] for vs in slide_videos(real) for v in vs}
        except (OSError, KeyError, ET.ParseError, zipfile.BadZipFile):
            allowed = set()
        if part not in allowed:
            return self.send(404, "not found", "text/plain")
        dest = media_part(real, part)
        self.raw(dest, mimetypes.guess_type(part)[0] or "application/octet-stream")

    def page_png(self, real, n):
        try:
            pages = page_images(real)
            png = pages[int(n) - 1]
        except (ValueError, IndexError, OSError, subprocess.SubprocessError):
            return self.send(404, "no such page", "text/plain")
        with open(png, "rb") as f:
            self.send(200, f.read(), "image/png", {"Cache-Control": "private, max-age=86400"})

    def entries(self, real):
        out = []
        for e in os.scandir(real):
            if e.name in TREE_HIDE or is_denied(os.path.join(real, e.name)):
                continue
            try:
                out.append((e, e.is_dir(), e.stat()))
            except OSError:
                continue
        return sorted(out, key=lambda t: (not t[1], natural(t[0].name)))

    def ls_json(self, real):
        try:
            entries = self.entries(real)
        except PermissionError:
            return self.send_json(403, {"entries": []})
        body = {"entries": [{"n": e.name, "d": is_dir} for e, is_dir, _ in entries[:TREE_MAX]]}
        if len(entries) > TREE_MAX:
            body["more"] = len(entries) - TREE_MAX
        self.send_json(200, body)

    def listing(self, path, real):
        try:
            entries = self.entries(real)
        except PermissionError:
            return self.respond(403, path, empty_state("lock", "Permission denied", "The server cannot list this folder."),
                                kind="empty")
        rows = []
        for e, is_dir, st in entries:
            href = urllib.parse.quote(e.name) + ("/" if is_dir else "") + "?view"
            cls = "dir" if is_dir else file_type(e.name)
            size = "" if is_dir else human(st.st_size)
            rows.append(f'<a class="frow" href="{href}"><span class="fname"><span class="fi {cls}"></span>'
                        f'<span class="nm">{html.escape(e.name)}</span></span><span class="fsize">{size}</span>'
                        f'<span class="ftime" data-t="{int(st.st_mtime)}"></span></a>')
        if not rows:
            rows.append('<div class="fnone">This folder is empty.</div>')
        body = ('<div class="files"><div class="fhead"><span>Name</span><span class="fsize">Size</span>'
                f'<span class="ftime">Modified</span></div>{"".join(rows)}</div>')
        readme = next((os.path.join(real, e.name) for e, is_dir, _ in entries
                       if not is_dir and e.name.lower() in ("readme.md", "context.md")), None)
        if readme:
            text, _ = read_text(readme)
            if text:
                body += (f'<section class="readme"><div class="readme-head"><span class="fi t-md"></span>'
                         f'{html.escape(os.path.basename(readme))}</div><article class="md">{render_markdown(text)}</article></section>')
        self.respond(200, path, f'<div class="listing">{body}</div>', kind="folder",
                     info={"items": len(entries), "lang": "Folder"})

    def render(self, path, real, line, end, want_src, want_pages=False):
        st = os.stat(real)
        name = os.path.basename(real)
        ext = os.path.splitext(name)[1].lower()
        info = {"size": human(st.st_size), "mtime": int(st.st_mtime)}
        raw = {"label": "Raw", "href": "?raw", "icon": "file", "native": True}
        source = {"label": "Source", "href": "?view&src", "icon": "code"}
        rendered = {"label": "Preview", "href": "?view", "icon": "eye"}
        download = {"label": "Download", "href": "?raw", "icon": "download", "native": True, "download": name}

        if ext in IMAGE_EXTS:
            body = f'<div class="media"><img src="?raw" alt="{html.escape(name)}"></div>'
            return self.respond(200, path, body, "media", dict(info, lang="Image"), [raw])
        if ext in VIDEO_EXTS:
            body = f'<div class="media"><video src="?raw" controls playsinline preload="metadata"></video></div>'
            return self.respond(200, path, body, "media", dict(info, lang="Video"), [raw, download])
        if ext in (".html", ".htm") and not want_src:
            body = f'<iframe src="?embed" title="{html.escape(name)}"></iframe>'
            full = {"label": "Open page", "href": "?embed", "icon": "external", "native": True, "target": "_blank"}
            return self.respond(200, path, body, "frame", dict(info, lang="HTML"), [source, full])
        if ext in SLIDE_EXTS or ext == ".docx" and want_pages:
            what = "Slide" if ext in SLIDE_EXTS else "Page"
            reflow = {"label": "Text", "href": "?view", "icon": "eye"}
            return self.pages(path, real, info, what, ([] if ext in SLIDE_EXTS else [reflow]) + [download])
        if ext == ".docx" and not want_src:
            with open(real, "rb") as f:
                doc = mammoth.convert_to_html(f).value
            pages = {"label": "Pages", "href": "?view&pages", "icon": "file"}
            return self.respond(200, path, f'<article class="md">{doc}</article>', "doc",
                                dict(info, lang="Word document"), [pages, download])
        if ext in SHEET_EXTS:
            return self.workbook(path, real, info, [download])
        if ext in DOWNLOAD_EXTS:
            body = empty_state("file", name, f"{info['size']} · no preview for this file type.",
                               [("Download", "?raw", True, f' download="{html.escape(name)}" data-native')])
            return self.respond(200, path, body, "empty", info, [download])

        text, note = read_text(real)
        notes = f'<div class="callout">{html.escape(note)}</div>' if note else ""
        if text is None:
            body = empty_state("file", name, f"{info['size']} · binary file.",
                               [("Download", "?raw", True, f' download="{html.escape(name)}" data-native')])
            return self.respond(200, path, body, "empty", info, [download])
        info["lines"] = text.count("\n") + (0 if text.endswith("\n") else 1)

        if ext in (".md", ".markdown") and not want_src and line is None:
            body = notes + f'<article class="md">{render_markdown(text)}</article>'
            return self.respond(200, path, body, "doc", dict(info, lang="Markdown"), [source, raw])
        if ext in (".csv", ".tsv") and not want_src:
            return self.respond(200, path, notes + self.csv_table(text, ext), "table",
                                dict(info, lang="CSV" if ext == ".csv" else "TSV"), [source, raw])

        if len(text) > MAX_HIGHLIGHT:
            code = "".join(f'<span id="L-{i}"><span class="linenos">{i}</span>{html.escape(t)}\n</span>'
                           for i, t in enumerate(text.splitlines(), 1))
            body, lang = f'<div class="hl"><pre><code>{code}</code></pre></div>', "Plain text"
        else:
            lexer = lexer_for(real, text)
            fmt = HtmlFormatter(cssclass="hl", linenos="inline", linespans="L", wrapcode=True)
            body, lang = highlight(text, lexer, fmt), lexer.name
        actions = [raw]
        if ext in (".md", ".markdown", ".html", ".htm", ".csv", ".tsv"):
            actions = [rendered, raw]
        self.respond(200, path, notes + body, "code", dict(info, lang=lang), actions, line, end)

    def no_preview(self, path, name, info, actions, why):
        body = empty_state("file", name, f"{info['size']} · {html.escape(why)}",
                           [("Download", "?raw", True, f' download="{html.escape(name)}" data-native')])
        return self.respond(200, path, body, "empty", info, actions)

    def pages(self, path, real, info, what, actions):
        try:
            pngs = page_images(real)
        except (OSError, subprocess.SubprocessError) as e:
            err = getattr(e, "stderr", b"") or b""
            sys.stderr.write(f"convert failed: {real}: {e} {err.decode(errors='replace')[-500:]}\n")
            return self.no_preview(path, os.path.basename(real), info, actions, "conversion failed.")
        if not pngs:
            return self.no_preview(path, os.path.basename(real), info, actions, "nothing to show.")
        v = info["mtime"]
        videos = []
        if real.lower().endswith(".pptx"):
            try:
                videos = slide_videos(real)
            except (OSError, KeyError, ET.ParseError, zipfile.BadZipFile) as e:
                sys.stderr.write(f"slide videos failed: {real}: {e}\n")
            if len(videos) != len(pngs):  # cannot line slides up with pages
                videos = []
        figs = []
        for i, png in enumerate(pngs, 1):
            w, h = png_size(png)
            clips = "".join(
                f'<video src="?media={urllib.parse.quote(c["part"])}&amp;v={v}" controls playsinline preload="metadata" '
                f'style="left:{c["x"]:.3f}%;top:{c["y"]:.3f}%;width:{c["w"]:.3f}%;height:{c["h"]:.3f}%"></video>'
                for c in (videos[i - 1] if videos else []))
            figs.append(f'<figure id="{what.lower()}-{i}"><div class="slide"><img src="?page={i}&amp;v={v}" width="{w}" height="{h}" '
                        f'loading="{"eager" if i <= 2 else "lazy"}" alt="{what} {i}">{clips}</div>'
                        f'<figcaption>{what} {i} / {len(pngs)}</figcaption></figure>')
        kind = "slides" if what == "Slide" else "slides paged"
        lang = "Presentation" if what == "Slide" else "Word document"
        return self.respond(200, path, f'<div class="pages">{"".join(figs)}</div>', kind,
                            dict(info, lang=f"{lang} · {len(pngs)} {what.lower()}{'s' * (len(pngs) != 1)}"), actions)

    def workbook(self, path, real, info, actions):
        try:
            wb = openpyxl.load_workbook(real, data_only=True)
        except Exception as e:  # openpyxl raises a zoo of types on odd files
            sys.stderr.write(f"xlsx failed: {real}: {e!r}\n")
            return self.no_preview(path, os.path.basename(real), info, actions, "could not read this workbook.")
        sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"] or wb.worksheets
        tabs, panes = [], []
        for i, ws in enumerate(sheets):
            table, note = sheet_table(ws)
            note = f'<div class="callout">{html.escape(note)}</div>' if note else ""
            tabs.append(f'<button class="sheet-tab{" on" if i == 0 else ""}" data-sheet="{i}">{html.escape(ws.title)}</button>')
            panes.append(f'<section class="sheet" data-sheet="{i}"{" hidden" if i else ""}>{note}{table}</section>')
        wb.close()
        body = f'<div class="book">{"".join(panes)}<nav class="sheet-tabs">{"".join(tabs)}</nav></div>'
        return self.respond(200, path, body, "table", dict(info, lang=f"Spreadsheet · {len(sheets)} sheet{'s' * (len(sheets) != 1)}"), actions)

    def csv_table(self, text, ext):
        reader = csv.reader(io.StringIO(text), delimiter="\t" if ext == ".tsv" else ",")
        rows = []
        for i, row in enumerate(reader):
            if i > CSV_ROWS:
                break
            tag = "th" if i == 0 else "td"
            rows.append("<tr>" + "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in row) + "</tr>")
        more = f'<div class="callout">Showing the first {CSV_ROWS} rows.</div>' if len(rows) > CSV_ROWS else ""
        return f'{more}<div class="grid"><table>{"".join(rows)}</table></div>'


if __name__ == "__main__":
    ThreadingHTTPServer.daemon_threads = True
    print(f"fileview on http://{HOST}:{PORT}", flush=True)
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
